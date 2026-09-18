#!/usr/bin/env python3
"""
Export HTML visual reports to styled PowerPoint (.pptx) presentations.

Usage:
    python export_pptx.py report.html output.pptx

Dependencies:
    pip install python-pptx beautifulsoup4 lxml
"""

import os
import sys
from pathlib import Path
from typing import Iterable, Optional

try:
    from bs4 import BeautifulSoup, Tag
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    print("Missing dependencies. Install with:")
    print("  pip install python-pptx beautifulsoup4 lxml")
    sys.exit(1)


SLIDE_W = 13.333
SLIDE_H = 7.5

ACCENTS = {
    "cyan": (34, 211, 238),
    "violet": (168, 85, 247),
    "amber": (249, 115, 22),
    "emerald": (16, 185, 129),
    "rose": (244, 63, 94),
    "blue": (59, 130, 246),
    "teal": (20, 184, 166),
    "orange": (249, 115, 22),
    "purple": (139, 92, 246),
    "green": (16, 185, 129),
}

THEMES = {
    "summit-dark": {
        "bg": (10, 20, 37),
        "panel": (23, 36, 58),
        "panel_alt": (30, 45, 68),
        "text": (248, 250, 252),
        "muted": (203, 213, 225),
        "soft": (148, 163, 184),
        "accent": "cyan",
        "secondary": "amber",
    },
    "signal-grid": {
        "bg": (15, 23, 42),
        "panel": (24, 24, 40),
        "panel_alt": (34, 32, 62),
        "text": (248, 250, 252),
        "muted": (191, 219, 254),
        "soft": (148, 163, 184),
        "accent": "violet",
        "secondary": "cyan",
    },
    "editorial-stage": {
        "bg": (17, 24, 39),
        "panel": (31, 41, 55),
        "panel_alt": (55, 65, 81),
        "text": (250, 250, 249),
        "muted": (226, 232, 240),
        "soft": (148, 163, 184),
        "accent": "rose",
        "secondary": "amber",
    },
    "ember-premium": {
        "bg": (24, 24, 27),
        "panel": (50, 33, 28),
        "panel_alt": (68, 43, 34),
        "text": (250, 250, 249),
        "muted": (231, 229, 228),
        "soft": (168, 162, 158),
        "accent": "amber",
        "secondary": "rose",
    },
}


def rgb(value: tuple[int, int, int]) -> RGBColor:
    return RGBColor(*value)


def mix(color_a: tuple[int, int, int], color_b: tuple[int, int, int], ratio: float) -> tuple[int, int, int]:
    return tuple(int(color_a[i] * ratio + color_b[i] * (1 - ratio)) for i in range(3))


def extract_text(element: Optional[Tag], separator: str = " ") -> str:
    if not element:
        return ""
    return " ".join(part.strip() for part in element.get_text(separator=separator).split() if part.strip())


def class_string(element: Optional[Tag]) -> str:
    if not element:
        return ""
    classes = element.get("class", [])
    return " ".join(classes) if isinstance(classes, list) else str(classes)


def has_class_fragment(element: Optional[Tag], fragment: str) -> bool:
    return fragment in class_string(element)


def theme_for_name(name: Optional[str]) -> dict:
    return THEMES.get(name or "", THEMES["summit-dark"])


def accent_for_name(name: Optional[str], theme: dict) -> tuple[int, int, int]:
    return ACCENTS.get(name or "", ACCENTS[theme["accent"]])


def parse_accent(element: Optional[Tag], theme: dict) -> tuple[int, int, int]:
    if element and element.get("data-accent"):
        return accent_for_name(element.get("data-accent"), theme)

    cls = class_string(element)
    for key, value in ACCENTS.items():
        if key in cls:
            return value
    return ACCENTS[theme["accent"]]


def first_heading(node: Tag) -> str:
    for tag_name in ("h1", "h2", "h3", "h4"):
        element = node.find(tag_name)
        text = extract_text(element)
        if text:
            return text
    return ""


def summary_text(node: Tag) -> str:
    for selector in (".ppt-subtitle", ".ppt-summary", ".ppt-note"):
        element = node.select_one(selector)
        text = extract_text(element)
        if text:
            return text

    for paragraph in node.find_all("p"):
        if not ancestor_has_fragment(paragraph, ("ppt-panel", "ppt-metric-card")):
            text = extract_text(paragraph)
            if text:
                return text
    return ""


def ancestor_has_fragment(element: Tag, fragments: Iterable[str]) -> bool:
    for parent in element.parents:
        cls = class_string(parent)
        if any(fragment in cls for fragment in fragments):
            return True
    return False


def find_text_by_class(node: Tag, fragments: tuple[str, ...]) -> str:
    for candidate in node.find_all(True):
        cls = class_string(candidate)
        if any(fragment in cls for fragment in fragments):
            text = extract_text(candidate)
            if text:
                return text
    return ""


def parse_metric_cards(node: Tag, theme: dict) -> list[dict]:
    cards = node.select(".ppt-metric-card")
    if not cards:
        cards = [
            div
            for div in node.find_all("div")
            if "rounded" in class_string(div)
            and ("gradient" in class_string(div) or "bg-gradient" in class_string(div))
        ]

    metrics = []
    for card in cards[:4]:
        label = find_text_by_class(card, ("ppt-metric-label",)) or find_text_by_class(card, ("text-sm",))
        value = find_text_by_class(card, ("ppt-metric-value", "text-4xl", "text-3xl", "text-5xl"))
        note = find_text_by_class(card, ("ppt-metric-note",))

        if not label:
            texts = [extract_text(div) for div in card.find_all(["div", "span", "p"]) if extract_text(div)]
            if texts:
                label = texts[0]
            if len(texts) > 2 and not note:
                note = texts[-1]

        if not value:
            texts = [extract_text(div) for div in card.find_all(["div", "span", "p"]) if extract_text(div)]
            value = next((text for text in texts if any(char.isdigit() for char in text)), "")

        if label or value:
            metrics.append(
                {
                    "label": label,
                    "value": value,
                    "note": note,
                    "accent": parse_accent(card, theme),
                }
            )
    return metrics


def parse_panels(node: Tag) -> list[dict]:
    panels = node.select(".ppt-panel")
    if not panels:
        panels = []
        for div in node.find_all("div", recursive=False):
            if div.find(["h3", "h4"]):
                panels.append(div)

    parsed = []
    for panel in panels[:4]:
        title = first_heading(panel)
        bullets = [extract_text(li) for li in panel.find_all("li") if extract_text(li)]
        paragraphs = [
            extract_text(p)
            for p in panel.find_all("p")
            if extract_text(p) and not ancestor_has_fragment(p, ("ppt-metric-card",))
        ]
        parsed.append({"title": title, "bullets": bullets[:5], "paragraphs": paragraphs[:2]})
    return parsed


def parse_bullets(node: Tag) -> list[str]:
    bullets = [extract_text(li) for li in node.find_all("li") if extract_text(li)]
    if bullets:
        return bullets[:5]

    paragraphs = []
    for paragraph in node.find_all("p"):
        if ancestor_has_fragment(paragraph, ("ppt-panel", "ppt-metric-card")):
            continue
        text = extract_text(paragraph)
        if text:
            paragraphs.append(text)
    return paragraphs[:4]


def parse_table(node: Tag) -> tuple[list[str], list[list[str]]]:
    table = node.find("table")
    if not table:
        return [], []

    headers = [extract_text(th) for th in table.find_all("th")]
    rows = []
    for tr in table.find_all("tr")[1:]:
        row = [extract_text(td) for td in tr.find_all(["td", "th"])]
        if row:
            rows.append(row)
    return headers, rows[:6]


def local_images(node: Tag, html_path: str) -> list[Path]:
    resolved = []
    for image in node.find_all("img"):
        src = image.get("src", "")
        if src and not src.startswith("http"):
            candidate = (Path(html_path).parent / src).resolve()
            if candidate.exists():
                resolved.append(candidate)
    return resolved


def add_textbox(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    *,
    size: int = 18,
    bold: bool = False,
    color: tuple[int, int, int] = (255, 255, 255),
    align: PP_ALIGN = PP_ALIGN.LEFT,
    font_name: Optional[str] = None,
    line_spacing: float = 1.25,
    vertical: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = vertical
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    font = paragraph.font
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = rgb(color)
    if font_name:
        font.name = font_name


def add_paragraphs(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    items: list[str],
    *,
    size: int,
    color: tuple[int, int, int],
    bullet: bool = False,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.clear()

    first = True
    for item in items:
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        paragraph.text = f"• {item}" if bullet else item
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.line_spacing = 1.25
        font = paragraph.font
        font.size = Pt(size)
        font.color.rgb = rgb(color)


def add_panel_shape(slide, left: float, top: float, width: float, height: float, fill_rgb: tuple[int, int, int], line_rgb: tuple[int, int, int]) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill_rgb)
    shape.fill.transparency = 0.12
    shape.line.color.rgb = rgb(line_rgb)
    shape.line.transparency = 0.35


def style_slide(slide, theme: dict, accent_rgb: tuple[int, int, int], index: int, total: int) -> None:
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(theme["bg"])

    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.35),
        Inches(0.45),
        Inches(0.08),
        Inches(1.15),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(accent_rgb)
    bar.line.fill.background()

    glow_primary = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(10.7),
        Inches(-0.7),
        Inches(3.2),
        Inches(3.2),
    )
    glow_primary.fill.solid()
    glow_primary.fill.fore_color.rgb = rgb(mix(accent_rgb, theme["bg"], 0.7))
    glow_primary.fill.transparency = 0.68
    glow_primary.line.fill.background()

    secondary_rgb = accent_for_name(theme["secondary"], theme)
    glow_secondary = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(-0.6),
        Inches(5.3),
        Inches(2.6),
        Inches(2.6),
    )
    glow_secondary.fill.solid()
    glow_secondary.fill.fore_color.rgb = rgb(mix(secondary_rgb, theme["bg"], 0.65))
    glow_secondary.fill.transparency = 0.78
    glow_secondary.line.fill.background()

    border = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.18),
        Inches(0.18),
        Inches(12.95),
        Inches(7.14),
    )
    border.fill.background()
    border.line.color.rgb = rgb(mix(theme["panel_alt"], theme["text"], 0.25))
    border.line.transparency = 0.75

    add_textbox(
        slide,
        12.25,
        6.95,
        0.65,
        0.25,
        f"{index}/{total}",
        size=10,
        color=theme["soft"],
        align=PP_ALIGN.RIGHT,
    )


def infer_slide_type(node: Tag) -> str:
    explicit = node.get("data-slide-type")
    if explicit:
        return explicit

    if node.find("h1"):
        return "cover"
    if node.find("table"):
        return "table"
    if node.select(".ppt-metric-card"):
        return "metrics"
    if node.select(".ppt-panel"):
        return "two-column"
    if node.find_all("li"):
        return "bullets"
    return "content"


def render_cover_slide(prs: Presentation, node: Tag, theme: dict, index: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent_rgb = parse_accent(node, theme)
    style_slide(slide, theme, accent_rgb, index, total)

    kicker = extract_text(node.select_one(".ppt-kicker")) or "Presentation"
    title = first_heading(node) or "Report"
    subtitle = summary_text(node)

    add_textbox(slide, 0.72, 0.8, 5.8, 0.35, kicker.upper(), size=12, color=accent_rgb)
    add_textbox(slide, 0.72, 1.35, 7.2, 2.2, title, size=29, bold=True, color=theme["text"])
    if subtitle:
        add_textbox(slide, 0.72, 3.55, 6.3, 1.5, subtitle, size=17, color=theme["muted"], line_spacing=1.35)

    panels = parse_panels(node)
    if panels:
        panel = panels[0]
        add_panel_shape(slide, 8.05, 1.55, 4.35, 3.1, theme["panel"], mix(accent_rgb, theme["text"], 0.35))
        if panel["title"]:
            add_textbox(slide, 8.35, 1.85, 3.7, 0.45, panel["title"], size=16, bold=True, color=theme["text"])
        panel_items = panel["bullets"] or panel["paragraphs"]
        if panel_items:
            add_paragraphs(slide, 8.35, 2.4, 3.6, 1.9, panel_items[:4], size=13, color=theme["muted"], bullet=bool(panel["bullets"]))
    else:
        note = summary_text(node)
        if note:
            add_panel_shape(slide, 8.05, 1.85, 4.35, 2.6, theme["panel"], mix(accent_rgb, theme["text"], 0.35))
            add_textbox(slide, 8.35, 2.15, 3.7, 1.8, note, size=14, color=theme["muted"], line_spacing=1.3)


def render_section_slide(prs: Presentation, node: Tag, theme: dict, index: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent_rgb = parse_accent(node, theme)
    style_slide(slide, theme, accent_rgb, index, total)

    kicker = extract_text(node.select_one(".ppt-kicker")) or "Section"
    title = first_heading(node)
    subtitle = summary_text(node)

    add_textbox(slide, 0.9, 1.2, 4.5, 0.35, kicker.upper(), size=12, color=accent_rgb)
    add_textbox(slide, 0.9, 2.1, 8.2, 1.5, title, size=31, bold=True, color=theme["text"])
    if subtitle:
        add_textbox(slide, 0.9, 4.05, 5.8, 1.1, subtitle, size=18, color=theme["muted"], line_spacing=1.35)


def render_metrics_slide(prs: Presentation, node: Tag, theme: dict, index: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent_rgb = parse_accent(node, theme)
    style_slide(slide, theme, accent_rgb, index, total)

    title = first_heading(node) or "关键指标"
    summary = summary_text(node)
    metrics = parse_metric_cards(node, theme)

    add_textbox(slide, 0.72, 0.78, 7.0, 0.55, title, size=24, bold=True, color=theme["text"])
    if summary:
        add_textbox(slide, 0.72, 1.38, 8.2, 0.6, summary, size=14, color=theme["muted"])

    count = max(1, min(len(metrics), 4))
    if count == 4:
        positions = [(0.72, 2.1), (6.72, 2.1), (0.72, 4.35), (6.72, 4.35)]
        size = (5.2, 1.75)
    elif count == 3:
        positions = [(0.72, 2.55), (4.45, 2.55), (8.18, 2.55)]
        size = (3.45, 2.1)
    elif count == 2:
        positions = [(0.95, 2.65), (6.7, 2.65)]
        size = (5.3, 2.2)
    else:
        positions = [(1.3, 2.55)]
        size = (10.6, 2.35)

    for metric, (left, top) in zip(metrics or [{"label": "", "value": "", "note": "", "accent": accent_rgb}], positions):
        width, height = size
        card_fill = mix(metric["accent"], theme["bg"], 0.72)
        add_panel_shape(slide, left, top, width, height, card_fill, mix(metric["accent"], theme["text"], 0.3))
        if metric["label"]:
            add_textbox(slide, left + 0.28, top + 0.16, width - 0.5, 0.32, metric["label"], size=11, color=theme["muted"])
        if metric["value"]:
            add_textbox(slide, left + 0.28, top + 0.5, width - 0.5, 0.62, metric["value"], size=22, bold=True, color=theme["text"])
        if metric["note"]:
            add_textbox(slide, left + 0.28, top + 1.15, width - 0.5, 0.5, metric["note"], size=11, color=theme["muted"])


def render_two_column_slide(prs: Presentation, node: Tag, theme: dict, index: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent_rgb = parse_accent(node, theme)
    style_slide(slide, theme, accent_rgb, index, total)

    title = first_heading(node)
    summary = summary_text(node)
    panels = parse_panels(node)

    add_textbox(slide, 0.72, 0.78, 8.0, 0.55, title, size=24, bold=True, color=theme["text"])
    if summary:
        add_textbox(slide, 0.72, 1.38, 8.0, 0.55, summary, size=14, color=theme["muted"])

    if not panels:
        bullets = parse_bullets(node)
        panels = [{"title": "", "bullets": bullets, "paragraphs": []}]

    positions = [(0.72, 2.1), (6.7, 2.1), (0.72, 4.5), (6.7, 4.5)]
    widths = [5.45, 5.45, 5.45, 5.45]
    heights = [2.0, 2.0, 1.7, 1.7]

    for idx, panel in enumerate(panels[:4]):
        left, top = positions[idx]
        width = widths[idx]
        height = heights[idx]
        add_panel_shape(slide, left, top, width, height, theme["panel"], mix(accent_rgb, theme["text"], 0.25))
        if panel["title"]:
            add_textbox(slide, left + 0.25, top + 0.18, width - 0.5, 0.36, panel["title"], size=16, bold=True, color=theme["text"])
        items = panel["bullets"] or panel["paragraphs"]
        if items:
            add_paragraphs(
                slide,
                left + 0.25,
                top + 0.55,
                width - 0.5,
                height - 0.65,
                items,
                size=12,
                color=theme["muted"],
                bullet=bool(panel["bullets"]),
            )


def render_bullets_slide(prs: Presentation, node: Tag, theme: dict, index: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent_rgb = parse_accent(node, theme)
    style_slide(slide, theme, accent_rgb, index, total)

    title = first_heading(node)
    summary = summary_text(node)
    bullets = parse_bullets(node)

    add_textbox(slide, 0.72, 0.78, 8.3, 0.55, title, size=24, bold=True, color=theme["text"])
    if summary:
        add_textbox(slide, 0.72, 1.38, 8.0, 0.55, summary, size=14, color=theme["muted"])

    add_panel_shape(slide, 0.72, 2.0, 7.2, 3.8, theme["panel"], mix(accent_rgb, theme["text"], 0.25))
    if bullets:
        add_paragraphs(slide, 1.0, 2.35, 6.6, 3.1, bullets, size=14, color=theme["muted"], bullet=True)


def render_table_slide(prs: Presentation, node: Tag, theme: dict, index: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent_rgb = parse_accent(node, theme)
    style_slide(slide, theme, accent_rgb, index, total)

    title = first_heading(node) or "数据表"
    summary = summary_text(node)
    headers, rows = parse_table(node)

    add_textbox(slide, 0.72, 0.78, 8.5, 0.55, title, size=24, bold=True, color=theme["text"])
    if summary:
        add_textbox(slide, 0.72, 1.38, 8.2, 0.55, summary, size=14, color=theme["muted"])

    if not headers or not rows:
        return

    total_rows = len(rows) + 1
    table_shape = slide.shapes.add_table(total_rows, len(headers), Inches(0.72), Inches(2.0), Inches(11.85), Inches(0.58 * total_rows))
    table = table_shape.table

    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(mix(accent_rgb, theme["bg"], 0.78))
        para = cell.text_frame.paragraphs[0]
        para.font.size = Pt(13)
        para.font.bold = True
        para.font.color.rgb = rgb(theme["text"])
        cell.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    for row_idx, row in enumerate(rows, start=1):
        fill_rgb = theme["panel"] if row_idx % 2 else theme["panel_alt"]
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill_rgb)
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(11)
            para.font.color.rgb = rgb(theme["muted"])
            cell.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE


def render_content_slide(prs: Presentation, node: Tag, theme: dict, index: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent_rgb = parse_accent(node, theme)
    style_slide(slide, theme, accent_rgb, index, total)

    title = first_heading(node) or "内容页"
    summary = summary_text(node)
    bullets = parse_bullets(node)

    add_textbox(slide, 0.72, 0.78, 8.5, 0.55, title, size=24, bold=True, color=theme["text"])
    if summary and summary not in bullets:
        add_textbox(slide, 0.72, 1.38, 8.2, 0.55, summary, size=14, color=theme["muted"])

    add_panel_shape(slide, 0.72, 2.0, 11.7, 3.95, theme["panel"], mix(accent_rgb, theme["text"], 0.25))
    content_items = bullets or ([summary] if summary else [])
    if content_items:
        add_paragraphs(
            slide,
            1.0,
            2.35,
            11.0,
            3.2,
            content_items[:5],
            size=14,
            color=theme["muted"],
            bullet=bool(bullets),
        )


def render_closing_slide(prs: Presentation, node: Tag, theme: dict, index: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent_rgb = parse_accent(node, theme)
    style_slide(slide, theme, accent_rgb, index, total)

    title = first_heading(node) or "总结"
    bullets = parse_bullets(node)
    summary = summary_text(node)

    add_textbox(slide, 0.9, 1.1, 7.8, 1.0, title, size=30, bold=True, color=theme["text"])
    if summary:
        add_textbox(slide, 0.9, 2.3, 6.8, 0.9, summary, size=16, color=theme["muted"], line_spacing=1.35)

    add_panel_shape(slide, 7.8, 1.55, 4.2, 3.15, theme["panel"], mix(accent_rgb, theme["text"], 0.25))
    closing_points = bullets[:3] or ([summary] if summary else [])
    if closing_points:
        add_paragraphs(slide, 8.1, 1.95, 3.55, 2.45, closing_points, size=14, color=theme["muted"], bullet=False)


def render_image_slide(prs: Presentation, image_path: Path, caption: str, theme: dict, index: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent_rgb = ACCENTS[theme["accent"]]
    style_slide(slide, theme, accent_rgb, index, total)

    try:
        slide.shapes.add_picture(str(image_path), Inches(0.9), Inches(0.9), width=Inches(11.5))
    except Exception as exc:
        print(f"Warning: Could not add image {image_path}: {exc}")
        return

    if caption:
        add_textbox(slide, 0.9, 6.55, 11.5, 0.35, caption, size=11, color=theme["muted"], align=PP_ALIGN.CENTER)


def render_semantic_slide(prs: Presentation, node: Tag, theme: dict, index: int, total: int) -> None:
    slide_type = infer_slide_type(node)
    if slide_type == "cover":
        render_cover_slide(prs, node, theme, index, total)
    elif slide_type == "section":
        render_section_slide(prs, node, theme, index, total)
    elif slide_type == "metrics":
        render_metrics_slide(prs, node, theme, index, total)
    elif slide_type == "two-column":
        render_two_column_slide(prs, node, theme, index, total)
    elif slide_type == "bullets":
        render_bullets_slide(prs, node, theme, index, total)
    elif slide_type == "table":
        render_table_slide(prs, node, theme, index, total)
    elif slide_type == "closing":
        render_closing_slide(prs, node, theme, index, total)
    else:
        render_content_slide(prs, node, theme, index, total)


def render_legacy_dom(prs: Presentation, soup: BeautifulSoup, theme: dict, html_path: str) -> None:
    title = extract_text(soup.find("h1")) or "Report"
    subtitle = extract_text(soup.find("p"))

    cover_markup = BeautifulSoup(
        f"""
        <section data-slide-type="cover" data-accent="{theme['accent']}">
          <div class="ppt-kicker">Report</div>
          <h1>{title}</h1>
          <p class="ppt-subtitle">{subtitle}</p>
        </section>
        """,
        "lxml",
    ).section

    legacy_nodes: list[Tag] = [cover_markup]

    metric_cards = parse_metric_cards(soup, theme)
    if metric_cards:
        metric_markup = BeautifulSoup(
            '<section data-slide-type="metrics"><h2>关键指标</h2><p class="ppt-summary"></p><div class="ppt-metric-grid"></div></section>',
            "lxml",
        ).section
        grid = metric_markup.find("div")
        for metric in metric_cards[:4]:
            card = BeautifulSoup(
                f"""
                <div class="ppt-metric-card" data-accent="{theme['accent']}">
                  <div class="ppt-metric-label">{metric['label']}</div>
                  <div class="ppt-metric-value">{metric['value']}</div>
                  <div class="ppt-metric-note">{metric['note']}</div>
                </div>
                """,
                "lxml",
            ).div
            grid.append(card)
        legacy_nodes.append(metric_markup)

    content_cards = soup.find_all("div", class_=lambda x: x and "bg-white" in x and "rounded-xl" in x)
    for card in content_cards:
        if first_heading(card):
            legacy_nodes.append(card)

    for table in soup.find_all("table"):
        table_wrapper = BeautifulSoup('<section data-slide-type="table"><h2>数据表</h2></section>', "lxml").section
        table_wrapper.append(table)
        legacy_nodes.append(table_wrapper)

    total = len(legacy_nodes)
    for index, node in enumerate(legacy_nodes, start=1):
        render_semantic_slide(prs, node, theme, index, total)

    base_index = len(legacy_nodes)
    for image in soup.find_all("img"):
        src = image.get("src", "")
        if src and not src.startswith("http"):
            image_path = (Path(html_path).parent / src).resolve()
            if image_path.exists():
                base_index += 1
                render_image_slide(prs, image_path, image.get("alt", ""), theme, base_index, base_index)


def html_to_pptx(html_path: str, output_path: str) -> None:
    with open(html_path, "r", encoding="utf-8") as file:
        soup = BeautifulSoup(file.read(), "lxml")

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    body = soup.body
    theme = theme_for_name(body.get("data-ppt-theme") if body else None)

    slide_nodes = soup.select("section.ppt-slide, div.ppt-slide, section.slide, div.slide")

    if slide_nodes:
        total = len(slide_nodes)
        for index, node in enumerate(slide_nodes, start=1):
            render_semantic_slide(prs, node, theme, index, total)
    else:
        render_legacy_dom(prs, soup, theme, html_path)

    prs.save(output_path)
    print(f"PowerPoint saved to: {output_path}")


def main() -> None:
    if len(sys.argv) < 3:
        print("Usage: python export_pptx.py <input.html> <output.pptx>")
        sys.exit(1)

    html_path = sys.argv[1]
    output_path = sys.argv[2]

    if not os.path.exists(html_path):
        print(f"Error: File not found: {html_path}")
        sys.exit(1)

    html_to_pptx(html_path, output_path)


if __name__ == "__main__":
    main()
